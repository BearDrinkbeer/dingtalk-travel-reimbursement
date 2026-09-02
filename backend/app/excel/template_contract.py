from __future__ import annotations

import posixpath
import re
from collections.abc import Iterable
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path, PurePosixPath
from typing import BinaryIO
from urllib.parse import urlsplit
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.core.errors import ApiError
from app.domain.categories import EXPENSE_CATEGORIES


@dataclass(frozen=True, slots=True)
class ExcelTemplateMapping:
    sheet_name: str = "费用报销模板"
    employee_name_cell: str = "C2"
    department_cell: str = "E2"
    project_cell: str = "G2"
    detail_start_row: int = 4
    detail_end_row: int = 54
    category_column: str = "B"
    date_column: str = "C"
    description_column: str = "D"
    amount_column: str = "G"
    currency_column: str = "H"
    receipt_count_column: str = "I"
    uppercase_amount_cell: str = "C55"
    total_amount_cell: str = "G55"
    total_receipt_count_cell: str = "I55"
    print_area: str = "'费用报销模板'!$B$1:$I$55"

    @property
    def detail_capacity(self) -> int:
        return self.detail_end_row - self.detail_start_row + 1

    @property
    def total_row(self) -> int:
        return self.detail_end_row + 1

    @property
    def print_end_row(self) -> int:
        return self.total_row

    def print_area_for(self, end_row: int) -> str:
        return f"'{self.sheet_name}'!$B$1:$I${end_row}"


@dataclass(frozen=True, slots=True)
class ExcelSheetLayout:
    detail_end_row: int
    total_row: int
    print_end_row: int

    @property
    def uppercase_amount_cell(self) -> str:
        return f"C{self.total_row}"

    @property
    def total_amount_cell(self) -> str:
        return f"G{self.total_row}"

    @property
    def total_receipt_count_cell(self) -> str:
        return f"I{self.total_row}"


EXCEL_TEMPLATE = ExcelTemplateMapping()
BASE_SHEET_LAYOUT = ExcelSheetLayout(
    detail_end_row=EXCEL_TEMPLATE.detail_end_row,
    total_row=EXCEL_TEMPLATE.total_row,
    print_end_row=EXCEL_TEMPLATE.print_end_row,
)


def output_layout_for_line_count(line_count: int) -> ExcelSheetLayout:
    if line_count < 0:
        raise ValueError("line_count must not be negative")
    extra_rows = max(0, line_count - EXCEL_TEMPLATE.detail_capacity)
    return ExcelSheetLayout(
        detail_end_row=EXCEL_TEMPLATE.detail_end_row + extra_rows,
        total_row=EXCEL_TEMPLATE.total_row + extra_rows,
        print_end_row=EXCEL_TEMPLATE.print_end_row + extra_rows,
    )

MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 200
MAX_ARCHIVE_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 200

_RELATIONSHIP_NAMESPACE = "http://schemas.openxmlformats.org/package/2006/relationships"
_WORKBOOK_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/worksheet",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/styles",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/theme",
        "http://purl.oclc.org/ooxml/officeDocument/relationships/sharedStrings",
    }
)
_UNSAFE_RELATIONSHIP_SUFFIXES = (
    "/attachedtemplate",
    "/connections",
    "/control",
    "/externallink",
    "/hyperlink",
    "/oleobject",
    "/package",
    "/querytable",
    "/vbaproject",
)
_UNSAFE_MEMBER_PREFIXES = (
    "customui/",
    "xl/activex/",
    "xl/ctrlprops/",
    "xl/dialogsheets/",
    "xl/embeddings/",
    "xl/externallinks/",
    "xl/macrosheets/",
    "xl/querytables/",
)
_UNSAFE_MEMBER_NAMES = frozenset(
    {
        "xl/connections.xml",
        "xl/vbaproject.bin",
    }
)
_URI_SCHEME = re.compile(r"^[A-Za-z][A-Za-z0-9+.-]*:")
_SAFE_MEMBER_NAMES = frozenset(
    {
        "[Content_Types].xml",
        "_rels/.rels",
        "docProps/app.xml",
        "docProps/core.xml",
        "xl/_rels/workbook.xml.rels",
        "xl/sharedStrings.xml",
        "xl/styles.xml",
        "xl/workbook.xml",
    }
)
_SAFE_MEMBER_PATTERNS = (
    re.compile(r"xl/theme/theme[1-9][0-9]*\.xml"),
    re.compile(r"xl/worksheets/sheet[1-9][0-9]*\.xml"),
)

EXPECTED_CATEGORY_VALIDATION_FORMULA = (
    '"' + ",".join(category.name for category in EXPENSE_CATEGORIES) + '"'
)

REQUIRED_FIXED_LABELS: dict[str, str] = {
    "B1": "费用报销单",
    "B2": "姓名",
    "D2": "部门",
    "F2": "报销项目",
    "B3": "费用类别",
    "C3": "发生日期",
    "D3": "说明",
    "G3": "金额",
    "H3": "币种",
    "I3": "票据张数",
}


def _required_labels(layout: ExcelSheetLayout) -> dict[str, str]:
    return REQUIRED_FIXED_LABELS | {
        f"B{layout.total_row}": "总合计",
        f"H{layout.total_row}": "人民币",
    }


def _required_merges(layout: ExcelSheetLayout) -> frozenset[str]:
    return frozenset(
        {
            "B1:I1",
            "G2:I2",
            "D3:F3",
            f"C{layout.total_row}:F{layout.total_row}",
            *(
                f"D{row}:F{row}"
                for row in range(EXCEL_TEMPLATE.detail_start_row, layout.detail_end_row + 1)
            ),
        }
    )


def apply_output_page_setup(
    worksheet: Worksheet,
    layout: ExcelSheetLayout = BASE_SHEET_LAYOUT,
) -> None:
    """Apply the official form's print styling after loading the safe template."""

    worksheet.print_area = f"B1:I{layout.print_end_row}"
    worksheet.sheet_properties.pageSetUpPr.fitToPage = None
    worksheet.page_setup.paperSize = "9"
    worksheet.page_setup.orientation = "portrait"
    worksheet.page_setup.scale = None
    worksheet.page_setup.fitToWidth = None
    worksheet.page_setup.fitToHeight = None
    worksheet.page_margins.left = 0.75
    worksheet.page_margins.right = 0.75
    worksheet.page_margins.top = 1.0
    worksheet.page_margins.bottom = 1.0
    worksheet.page_margins.header = 0.5
    worksheet.page_margins.footer = 0.5
    # The official workbook leaves gridline visibility at Excel's default.
    worksheet.sheet_view.showGridLines = None


def template_invalid() -> ApiError:
    return ApiError(
        "EXCEL_TEMPLATE_INVALID",
        "Excel 模板结构无效，请联系管理员",
        500,
    )


def _invalid_archive_member(name: str) -> bool:
    if not name or "\x00" in name or "\\" in name:
        return True
    path = PurePosixPath(name)
    return path.is_absolute() or any(part in {"", ".", ".."} for part in path.parts)


def _safe_archive_member(name: str) -> bool:
    return name in _SAFE_MEMBER_NAMES or any(
        pattern.fullmatch(name) for pattern in _SAFE_MEMBER_PATTERNS
    )


def _relationship_source(name: str) -> str:
    directory, filename = posixpath.split(name)
    if directory == "_rels" and filename == ".rels":
        return ""
    if not directory.endswith("/_rels") or not filename.endswith(".rels"):
        raise template_invalid()
    return posixpath.join(directory[: -len("/_rels")], filename[: -len(".rels")])


def _validate_relationships(name: str, content: bytes, members: set[str]) -> None:
    try:
        root = ElementTree.fromstring(content)
    except ElementTree.ParseError:
        raise template_invalid() from None

    source = _relationship_source(name)
    for relationship in root.findall(f"{{{_RELATIONSHIP_NAMESPACE}}}Relationship"):
        relationship_type = relationship.attrib.get("Type", "")
        target = relationship.attrib.get("Target", "")
        target_mode = relationship.attrib.get("TargetMode", "")
        if not relationship_type or not target:
            raise template_invalid()
        if target_mode.lower() == "external":
            raise template_invalid()
        if relationship_type.lower().endswith(_UNSAFE_RELATIONSHIP_SUFFIXES):
            raise template_invalid()
        if (
            name == "xl/_rels/workbook.xml.rels"
            and relationship_type not in _WORKBOOK_RELATIONSHIP_TYPES
        ):
            raise template_invalid()
        parsed = urlsplit(target)
        if parsed.scheme or parsed.netloc or _URI_SCHEME.match(target) or "\\" in target:
            raise template_invalid()
        if target.startswith("/"):
            resolved = posixpath.normpath(target.lstrip("/"))
        else:
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source), target))
        if resolved == ".." or resolved.startswith("../") or resolved not in members:
            raise template_invalid()


def validate_xlsx_archive(source: Path | bytes | bytearray | BinaryIO) -> None:
    """Validate package structure without extracting or trying to sanitize it."""

    try:
        if isinstance(source, Path):
            if source.stat().st_size > MAX_TEMPLATE_BYTES:
                raise template_invalid()
            archive_source: Path | BytesIO | BinaryIO = source
        elif isinstance(source, bytes | bytearray):
            if len(source) > MAX_TEMPLATE_BYTES:
                raise template_invalid()
            archive_source = BytesIO(source)
        else:
            archive_source = source

        with ZipFile(archive_source) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ARCHIVE_ENTRIES:
                raise template_invalid()
            names = [entry.filename for entry in entries]
            if len(names) != len(set(names)) or len(names) != len(
                {name.casefold() for name in names}
            ):
                raise template_invalid()
            if any(_invalid_archive_member(name) for name in names):
                raise template_invalid()
            if any(not _safe_archive_member(name) for name in names):
                raise template_invalid()
            worksheet_parts = [
                name for name in names if re.fullmatch(r"xl/worksheets/sheet[1-9][0-9]*\.xml", name)
            ]
            if len(worksheet_parts) != 1:
                raise template_invalid()
            members = set(names)
            if "xl/workbook.xml" not in members or "xl/_rels/workbook.xml.rels" not in members:
                raise template_invalid()

            total_uncompressed = 0
            for entry in entries:
                total_uncompressed += entry.file_size
                if entry.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise template_invalid()
                if entry.file_size and not entry.compress_size:
                    raise template_invalid()
                if (
                    entry.compress_size
                    and entry.file_size / entry.compress_size > MAX_COMPRESSION_RATIO
                ):
                    raise template_invalid()
            if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise template_invalid()

            lowered = {name.casefold() for name in names}
            if any(name.endswith(".bin") for name in lowered):
                raise template_invalid()
            if any(name in _UNSAFE_MEMBER_NAMES for name in lowered):
                raise template_invalid()
            if any(name.startswith(_UNSAFE_MEMBER_PREFIXES) for name in lowered):
                raise template_invalid()

            for name in names:
                if name.endswith(".rels"):
                    _validate_relationships(name, archive.read(name), members)
    except ApiError:
        raise
    except (BadZipFile, OSError, ValueError):
        raise template_invalid() from None


def _iter_cells(workbook: Workbook) -> Iterable[object]:
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            yield from row


def validate_workbook_safety(
    workbook: Workbook,
    layout: ExcelSheetLayout = BASE_SHEET_LAYOUT,
) -> Worksheet:
    """Reject active content and any workbook shape outside the fixed contract."""

    if workbook.sheetnames != [EXCEL_TEMPLATE.sheet_name]:
        raise template_invalid()
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if worksheet.sheet_state != "visible":
        raise template_invalid()
    if len(getattr(workbook, "_external_links", ())) != 0:
        raise template_invalid()
    if workbook.defined_names:
        # openpyxl maps the allowed built-in Print_Area to worksheet.print_area.
        # Any remaining defined name is outside the fixed template contract.
        raise template_invalid()
    if len(worksheet.conditional_formatting) != 0:
        raise template_invalid()
    if worksheet.tables:
        raise template_invalid()
    if getattr(worksheet, "_pivots", ()):  # openpyxl exposes no public pivot collection.
        raise template_invalid()
    if getattr(worksheet, "_charts", ()) or getattr(worksheet, "_images", ()):
        raise template_invalid()
    for cell in _iter_cells(workbook):
        value = cell.value
        if cell.data_type == "f" or (isinstance(value, str) and value.startswith("=")):
            raise template_invalid()
        if cell.hyperlink is not None:
            raise template_invalid()
        if cell.comment is not None:
            raise template_invalid()
    validate_template_sheet(worksheet, layout)
    return worksheet


def validate_template_sheet(
    worksheet: Worksheet,
    layout: ExcelSheetLayout = BASE_SHEET_LAYOUT,
) -> None:
    """Fail closed when the maintained template no longer matches the mapping."""

    actual_merges = {str(item) for item in worksheet.merged_cells.ranges}
    if actual_merges != _required_merges(layout):
        raise template_invalid()
    if any(
        worksheet[cell].value != expected
        for cell, expected in _required_labels(layout).items()
    ):
        raise template_invalid()
    if worksheet.print_area != EXCEL_TEMPLATE.print_area_for(layout.print_end_row):
        raise template_invalid()
    if worksheet.page_setup.orientation != "portrait":
        raise template_invalid()
    if worksheet.page_setup.fitToWidth is not None or worksheet.page_setup.fitToHeight is not None:
        raise template_invalid()
    if worksheet.print_title_rows is not None or worksheet.print_title_cols is not None:
        raise template_invalid()
    if worksheet.max_row < layout.total_row:
        raise template_invalid()

    validations = list(worksheet.data_validations.dataValidation)
    if len(validations) != 1:
        raise template_invalid()
    validation = validations[0]
    if (
        str(validation.sqref)
        != f"B{EXCEL_TEMPLATE.detail_start_row}:B{layout.detail_end_row}"
        or validation.type != "list"
        or validation.formula1 != EXPECTED_CATEGORY_VALIDATION_FORMULA
        or validation.formula2 not in {None, ""}
        or validation.operator is not None
        or validation.allowBlank is not False
        or validation.showDropDown is not False
        or validation.showInputMessage is not False
        or validation.showErrorMessage is not False
        or validation.promptTitle is not None
        or validation.prompt is not None
        or validation.errorTitle is not None
        or validation.error is not None
        or validation.errorStyle is not None
        or validation.imeMode is not None
    ):
        raise template_invalid()
