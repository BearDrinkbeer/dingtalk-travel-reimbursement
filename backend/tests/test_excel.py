from __future__ import annotations

import hashlib
from io import BytesIO
from pathlib import Path
from urllib.parse import unquote
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from conftest import mock_login
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table

from app.excel.template_contract import EXCEL_TEMPLATE, apply_output_page_setup
from app.services.excel_generator import XLSX_MEDIA_TYPE

TEMPLATE_PATH = Path(__file__).parents[1] / "app" / "templates" / "expense_template.xlsx"
CANONICAL_TEMPLATE_SHA256 = "39ba8327577e3ea109aa462b82c0fde265f6e7fa6c50fc2a824c6890326b27c5"


def trip() -> dict[str, object]:
    return {
        "tripType": "business",
        "startDate": "2026-06-30",
        "startTime": "09:00",
        "endDate": "2026-07-07",
        "endTime": "18:00",
    }


def item(
    *,
    category: str = "other",
    date: str = "2026-06-30",
    display_date: str | None = None,
    description: str = "测试费用",
    amount: str = "1.00",
    receipt_count: int = 1,
) -> dict[str, object]:
    return {
        "category": category,
        "date": date,
        "displayDate": display_date or date,
        "description": description,
        "amount": amount,
        "receiptCount": receipt_count,
    }


def manual_body(*, items: list[dict[str, object]] | None = None) -> dict[str, object]:
    return {
        "project": {"mode": "manual", "text": "P-001 示例项目"},
        "trip": trip(),
        "items": items or [],
    }


def xlsx_from_response(response):
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    return load_workbook(BytesIO(response.content), data_only=False, keep_links=True)


def assert_template_invalid(client_factory, template_path: Path) -> None:
    client = client_factory(auth_mock_enabled=True, excel_template_path=template_path)
    csrf = mock_login(client)["csrfToken"]
    response = client.post(
        "/api/excel/generate",
        json=manual_body(),
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 500
    assert response.json()["error"] == {
        "code": "EXCEL_TEMPLATE_INVALID",
        "message": "Excel 模板结构无效，请联系管理员",
    }
    assert str(template_path) not in response.text


def assert_archive_has_no_active_content(content: bytes) -> None:
    with ZipFile(BytesIO(content)) as archive:
        names = {name.casefold() for name in archive.namelist()}
        assert not any(name.endswith(".bin") for name in names)
        assert not any(name.startswith("xl/externallinks/") for name in names)
        assert not any(name.startswith("xl/querytables/") for name in names)
        assert "xl/connections.xml" not in names
        assert not any(name.startswith("xl/activex/") for name in names)
        assert not any(name.startswith("xl/drawings/") for name in names)
        assert not any(name.startswith("xl/charts/") for name in names)
        assert not any(name.startswith("xl/tables/") for name in names)
        assert not any(name.startswith("xl/pivottables/") for name in names)
        assert not any(name.startswith("xl/pivotcache/") for name in names)
        assert not any(name.startswith("xl/slicers/") for name in names)
        assert not any(name.startswith("xl/comments") for name in names)
        for name in archive.namelist():
            if name.endswith(".rels"):
                relationships = archive.read(name).lower()
                assert b'targetmode="external"' not in relationships
                assert b'/hyperlink"' not in relationships


def template_layout(worksheet) -> dict[str, object]:
    return {
        "merges": sorted(str(value) for value in worksheet.merged_cells.ranges),
        "widths": {column: worksheet.column_dimensions[column].width for column in "ABCDEFGHI"},
        "heights": {row: worksheet.row_dimensions[row].height for row in range(1, 57)},
        "styles": {
            cell: (
                worksheet[cell].font.name,
                worksheet[cell].font.sz,
                worksheet[cell].font.bold,
                worksheet[cell].fill.fill_type,
                worksheet[cell].fill.fgColor.rgb,
                worksheet[cell].number_format,
                worksheet[cell].alignment.horizontal,
                worksheet[cell].alignment.vertical,
                worksheet[cell].alignment.wrap_text,
                worksheet[cell].border.left.style,
                worksheet[cell].border.right.style,
                worksheet[cell].border.top.style,
                worksheet[cell].border.bottom.style,
            )
            for cell in ("C2", "E2", "G2", "B4", "C4", "D4", "G4", "H4", "I4", "C55", "I55")
        },
        "print_area": worksheet.print_area,
        "orientation": worksheet.page_setup.orientation,
        "fit_width": worksheet.page_setup.fitToWidth,
        "fit_height": worksheet.page_setup.fitToHeight,
        "margins": (
            worksheet.page_margins.left,
            worksheet.page_margins.right,
            worksheet.page_margins.top,
            worksheet.page_margins.bottom,
            worksheet.page_margins.header,
            worksheet.page_margins.footer,
        ),
        "validations": [
            (str(rule.sqref), rule.type, rule.formula1)
            for rule in worksheet.data_validations.dataValidation
        ],
    }


def test_template_matches_official_form_format() -> None:
    workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert {column: worksheet.column_dimensions[column].width for column in "ABCDEFGHI"} == {
        "A": 2.6640625,
        "B": 10.83203125,
        "C": 8.83203125,
        "D": 7.1640625,
        "E": 17.83203125,
        "F": 6.1640625,
        "G": 12.5,
        "H": 7.33203125,
        "I": 9.5,
    }
    assert worksheet.row_dimensions[1].height == pytest.approx(28)
    assert worksheet.row_dimensions[2].height == pytest.approx(35)
    assert worksheet.row_dimensions[3].height == pytest.approx(28)
    assert all(
        worksheet.row_dimensions[row].height == pytest.approx(30) for row in range(4, 55)
    )
    assert worksheet.row_dimensions[55].height == pytest.approx(28)
    assert worksheet["B1"].font.name == "微软雅黑"
    assert worksheet["B1"].font.sz == 14
    assert worksheet["B1"].font.bold is True
    assert worksheet["B2"].fill.fgColor.rgb == "FFBFBFBF"
    assert worksheet["B55"].fill.fgColor.rgb == "FFB5C6EA"
    assert worksheet["G4"].number_format == "General"
    assert worksheet["G54"].number_format == worksheet["G4"].number_format
    assert worksheet["I54"].number_format == worksheet["I4"].number_format
    assert worksheet["G55"].number_format.startswith("_ \\¥*")
    assert worksheet.sheet_view.showGridLines in {None, True}
    workbook.close()


def test_excel_endpoint_requires_auth_csrf_and_selected_department(client_factory) -> None:
    anonymous = client_factory(auth_mock_enabled=True)
    assert anonymous.post("/api/excel/generate", json=manual_body()).status_code == 401

    authenticated = client_factory(auth_mock_enabled=True)
    login = mock_login(authenticated)
    assert authenticated.post("/api/excel/generate", json=manual_body()).status_code == 403
    assert (
        authenticated.post(
            "/api/excel/generate",
            json=manual_body(),
            headers={"X-CSRF-Token": login["csrfToken"]},
        ).status_code
        == 200
    )

    multi_department = client_factory(
        auth_mock_enabled=True,
        auth_mock_departments="100:部门甲,200:部门乙",
    )
    multi_login = mock_login(multi_department)
    response = multi_department.post(
        "/api/excel/generate",
        json=manual_body(),
        headers={"X-CSRF-Token": multi_login["csrfToken"]},
    )
    assert response.status_code == 409
    assert response.json()["error"]["code"] == "DEPARTMENT_REQUIRED"


@pytest.mark.parametrize(
    "spoofed",
    [
        {"employee": {"name": "伪造用户"}},
        {"department": "伪造部门"},
        {"totals": {"totalAmount": "1.00"}},
        {"days": 1},
        {"subsidyTotal": "1.00"},
        {"uppercaseAmount": "壹元整"},
    ],
)
def test_excel_request_rejects_client_authority_fields(client_factory, spoofed) -> None:
    client = client_factory(auth_mock_enabled=True)
    csrf = mock_login(client)["csrfToken"]
    response = client.post(
        "/api/excel/generate",
        json=manual_body() | spoofed,
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 422
    assert response.json()["error"]["code"] == "VALIDATION_ERROR"


def test_excel_exact_cells_order_totals_and_layout_are_preserved(client_factory) -> None:
    before_hash = hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest()
    assert before_hash == CANONICAL_TEMPLATE_SHA256
    template_workbook = load_workbook(TEMPLATE_PATH, data_only=False)
    template_worksheet = template_workbook[EXCEL_TEMPLATE.sheet_name]
    apply_output_page_setup(template_worksheet)
    expected_layout = template_layout(template_worksheet)
    template_workbook.close()

    client = client_factory(auth_mock_enabled=True)
    csrf = mock_login(client)["csrfToken"]
    body = {
        "project": {"mode": "manual", "text": "P-001 示例项目"},
        "trip": trip(),
        "items": [
            item(
                category="local_transport",
                date="2026-07-06",
                display_date="7月1日、7月6日",
                description="酒店-项目-酒店",
                amount="44.89",
                receipt_count=4,
            ),
            item(
                category="rail_fare",
                date="2026-06-30",
                description="北京南-合肥南",
                amount="454.00",
            ),
            item(
                category="rail_fare",
                date="2026-07-07",
                description="合肥南-北京南",
                amount="473.50",
            ),
        ],
    }
    response = client.post(
        "/api/excel/generate",
        json=body,
        headers={"X-CSRF-Token": csrf},
    )
    assert_archive_has_no_active_content(response.content)
    workbook = xlsx_from_response(response)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]

    assert worksheet["C2"].value == "开发测试用户"
    assert worksheet["E2"].value == "测试部门"
    assert worksheet["G2"].value == "P-001 示例项目"
    assert [worksheet[f"B{row}"].value for row in range(4, 8)] == [
        "火车票",
        "市内交通费",
        "火车票",
        "出差补助",
    ]
    assert worksheet["C4"].value == "6月30日"
    assert worksheet["C5"].value == "7月1日、7月6日"
    assert worksheet["C6"].value == "7月7日"
    assert worksheet["C7"].value == "7月7日"
    assert worksheet["D7"].value == "6/30—7/7，共8天出差补助"
    assert [worksheet[f"H{row}"].value for row in range(4, 8)] == ["人民币"] * 4
    assert [worksheet[f"I{row}"].value for row in range(4, 8)] == [1, 4, 1, 0]
    assert str(worksheet["G4"].value) == "454"
    assert str(worksheet["G5"].value) == "44.89"
    assert str(worksheet["G6"].value) == "473.5"
    assert str(worksheet["G7"].value) == "800"
    assert worksheet["C55"].value == "壹仟柒佰柒拾贰元叁角玖分"
    assert str(worksheet["G55"].value) == "1772.39"
    assert worksheet["G55"].number_format.startswith("_ \\¥*")
    assert worksheet["I55"].value == 6
    assert template_layout(worksheet) == expected_layout
    assert all(not worksheet.row_dimensions[row].hidden for row in range(4, 8))
    assert all(worksheet.row_dimensions[row].hidden for row in range(8, 55))
    assert not any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row)
    assert not any(cell.hyperlink is not None for row in worksheet.iter_rows() for cell in row)
    assert not any(cell.comment is not None for row in worksheet.iter_rows() for cell in row)
    assert len(worksheet.conditional_formatting) == 0
    assert not worksheet.tables
    assert not worksheet._pivots
    assert not worksheet._charts
    assert not worksheet._images
    assert len(workbook._external_links) == 0
    workbook.close()
    assert hashlib.sha256(TEMPLATE_PATH.read_bytes()).hexdigest() == before_hash
    assert not list(client.app.state.settings.temp_dir.rglob("*.xlsx"))


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_formula_injection_is_saved_as_text_for_all_external_text_fields(
    client_factory, prefix: str
) -> None:
    client = client_factory(
        auth_mock_enabled=True,
        auth_mock_user_name=f"{prefix}NAME()",
        auth_mock_departments=f"100:{prefix}DEPT()",
    )
    csrf = mock_login(client)["csrfToken"]
    body = manual_body(
        items=[
            item(
                display_date=f"  {prefix}DATE()",
                description=f"{prefix}DESC()",
            )
        ]
    )
    body["project"] = {"mode": "manual", "text": f"{prefix}PROJECT()"}
    response = client.post(
        "/api/excel/generate",
        json=body,
        headers={"X-CSRF-Token": csrf},
    )
    workbook = xlsx_from_response(response)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    for coordinate in ("C2", "E2", "G2", "C4", "D4"):
        assert worksheet[coordinate].data_type == "s"
        assert worksheet[coordinate].value.startswith("'")
    assert not any(cell.data_type == "f" for row in worksheet.iter_rows() for cell in row)
    workbook.close()


def test_excel_dynamically_extends_rows_and_enforces_configured_safety_limit(
    client_factory,
) -> None:
    client = client_factory(auth_mock_enabled=True, expense_max_items=200)
    csrf = mock_login(client)["csrfToken"]
    headers = {"X-CSRF-Token": csrf}
    seventy_five = [item(description=f"费用 {index}") for index in range(75)]
    valid = client.post(
        "/api/excel/generate",
        json=manual_body(items=seventy_five),
        headers=headers,
    )
    workbook = xlsx_from_response(valid)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert worksheet["B78"].value == "其他"
    assert worksheet["D78"].value == "费用 74"
    assert worksheet["B79"].value == "出差补助"
    assert worksheet["B80"].value == "总合计"
    assert str(worksheet["G80"].value) == "875"
    assert worksheet["I80"].value == 75
    assert "D55:F55" in {str(value) for value in worksheet.merged_cells.ranges}
    assert "D79:F79" in {str(value) for value in worksheet.merged_cells.ranges}
    assert "C80:F80" in {str(value) for value in worksheet.merged_cells.ranges}
    assert str(worksheet.data_validations.dataValidation[0].sqref) == "B4:B79"
    assert worksheet.print_area == "'费用报销模板'!$B$1:$I$80"
    assert worksheet["G55"].number_format == worksheet["G54"].number_format
    assert worksheet["G80"].number_format.startswith("_ \\¥*")
    assert not any(worksheet.row_dimensions[row].hidden for row in range(4, 80))
    workbook.close()

    two_hundred = [item(description=f"费用 {index}") for index in range(200)]
    maximum = client.post(
        "/api/excel/generate",
        json=manual_body(items=two_hundred),
        headers=headers,
    )
    workbook = xlsx_from_response(maximum)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert worksheet["B203"].value == "其他"
    assert worksheet["B204"].value == "出差补助"
    assert worksheet["B205"].value == "总合计"
    assert worksheet["I205"].value == 200
    assert worksheet.print_area == "'费用报销模板'!$B$1:$I$205"
    workbook.close()

    overflow = client.post(
        "/api/excel/generate",
        json=manual_body(items=two_hundred + [item(description="第 201 条")]),
        headers=headers,
    )
    assert overflow.status_code == 422
    assert overflow.json()["error"]["code"] == "TOO_MANY_EXPENSE_LINES"
    assert "200" in overflow.json()["error"]["message"]


def test_excel_without_trip_has_no_subsidy_line(client_factory) -> None:
    client = client_factory(auth_mock_enabled=True)
    csrf = mock_login(client)["csrfToken"]
    body = manual_body(
        items=[item(description="非出差报销", amount="16.90")],
    )
    body["trip"] = None

    response = client.post(
        "/api/excel/generate",
        json=body,
        headers={"X-CSRF-Token": csrf},
    )
    workbook = xlsx_from_response(response)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    assert worksheet["B4"].value == "其他"
    assert worksheet["D4"].value == "非出差报销"
    assert str(worksheet["G4"].value) == "16.9"
    assert worksheet.row_dimensions[5].hidden
    assert str(worksheet["G55"].value) == "16.9"
    assert worksheet["I55"].value == 1
    workbook.close()


@pytest.mark.parametrize("damage", ["missing_sheet", "altered_merge"])
def test_invalid_template_returns_stable_error(client_factory, tmp_path: Path, damage: str) -> None:
    bad_template = tmp_path / f"{damage}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if damage == "missing_sheet":
        worksheet.title = "WrongSheet"
    else:
        worksheet.unmerge_cells("D4:F4")
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


def test_hidden_second_sheet_with_formula_is_rejected(client_factory, tmp_path: Path) -> None:
    bad_template = tmp_path / "hidden-formula-sheet.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook.create_sheet("Hidden")
    worksheet.sheet_state = "hidden"
    worksheet["A1"] = "=1+1"
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


def test_print_titles_are_rejected(client_factory, tmp_path: Path) -> None:
    bad_template = tmp_path / "print-titles.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    workbook[EXCEL_TEMPLATE.sheet_name].print_title_rows = "1:4"
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize("validation_type", ["custom", "list_formula"])
def test_formula_bearing_data_validation_is_rejected(
    client_factory, tmp_path: Path, validation_type: str
) -> None:
    bad_template = tmp_path / f"validation-{validation_type}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    validation = workbook[EXCEL_TEMPLATE.sheet_name].data_validations.dataValidation[0]
    if validation_type == "custom":
        validation.type = "custom"
        validation.formula1 = 'WEBSERVICE("https://example.invalid")'
    else:
        validation.formula1 = '=WEBSERVICE("https://example.invalid")'
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize("carrier", ["conditional_formatting", "table", "chart", "comment"])
def test_other_formula_or_active_content_carriers_are_rejected(
    client_factory, tmp_path: Path, carrier: str
) -> None:
    bad_template = tmp_path / f"carrier-{carrier}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if carrier == "conditional_formatting":
        worksheet.conditional_formatting.add(
            "G4",
            FormulaRule(formula=['WEBSERVICE("https://example.invalid")']),
        )
    elif carrier == "table":
        worksheet["A26"] = "列一"
        worksheet["B26"] = "列二"
        worksheet["A27"] = "值一"
        worksheet["B27"] = "值二"
        worksheet.add_table(Table(displayName="UnexpectedTable", ref="A26:B27"))
    elif carrier == "chart":
        chart = BarChart()
        chart.add_data(
            Reference(worksheet, min_col=7, min_row=3, max_row=4),
            titles_from_data=True,
        )
        worksheet.add_chart(chart, "K2")
    else:
        worksheet["D4"].comment = Comment("不应存在的批注", "attacker")
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize("damage", ["formula", "hyperlink", "defined_name"])
def test_template_active_content_is_rejected(client_factory, tmp_path: Path, damage: str) -> None:
    bad_template = tmp_path / f"{damage}.xlsx"
    workbook = load_workbook(TEMPLATE_PATH)
    worksheet = workbook[EXCEL_TEMPLATE.sheet_name]
    if damage == "formula":
        worksheet["G4"] = "=1+1"
    elif damage == "hyperlink":
        worksheet["D4"].hyperlink = "https://example.invalid/"
        worksheet["D4"] = "链接"
    else:
        workbook.defined_names.add(DefinedName("UnexpectedName", attr_text="'费用报销模板'!$B$2"))
    workbook.save(bad_template)
    workbook.close()

    assert_template_invalid(client_factory, bad_template)


@pytest.mark.parametrize(
    ("member_name", "content"),
    [
        ("xl/connections.xml", b"<connections/>"),
        ("xl/externalLinks/externalLink1.xml", b"<externalLink/>"),
        ("xl/queryTables/queryTable1.xml", b"<queryTable/>"),
        ("xl/vbaProject.bin", b"not-a-real-macro"),
    ],
)
def test_injected_active_ooxml_part_is_rejected(
    client_factory,
    tmp_path: Path,
    member_name: str,
    content: bytes,
) -> None:
    bad_template = tmp_path / member_name.replace("/", "-")
    bad_template.write_bytes(TEMPLATE_PATH.read_bytes())
    with ZipFile(bad_template, "a", compression=ZIP_DEFLATED) as archive:
        archive.writestr(member_name, content)

    assert_template_invalid(client_factory, bad_template)


def test_filename_is_sanitized_and_rfc5987_encoded(client_factory) -> None:
    client = client_factory(
        auth_mock_enabled=True,
        auth_mock_user_name='测/试:*?"<>|用户',
    )
    csrf = mock_login(client)["csrfToken"]
    body = manual_body()
    body["project"] = {"mode": "manual", "text": "预算/项目:一"}
    response = client.post(
        "/api/excel/generate",
        json=body,
        headers={"X-CSRF-Token": csrf},
    )
    assert response.status_code == 200
    header = response.headers["content-disposition"]
    assert "filename=expense-report.xlsx" in header
    encoded = header.split("filename*=UTF-8''", 1)[1]
    filename = unquote(encoded)
    assert filename.startswith("差旅费报销单-") and filename.endswith(".xlsx")
    assert not any(character in filename for character in '/\\:*?"<>|')
